"""منبعِ اصلاحِ تاریخی از خروجی‌های اصلیِ Joolio (فقط خواندن؛ هیچ معامله‌ای وارد نمی‌شود).

  python tools/extract_stage_source.py OUT.json FILE1.xlsx [FILE2.xlsx ...]

برای هر ردیف فقط: شمارهٔ معامله، مقدارِ خامِ «تغییر مرحله»، برنده‌بودن (همان canonFunnelِ ایمپورت)، هشِ نامِ مشتری
و منشأ (sha256ِ فایل، sheet، شمارهٔ ردیف). نامِ مشتری و دادهٔ شخصی ذخیره نمی‌شود. نیازمندی: openpyxl.
"""
import hashlib
import json
import os
import re
import sys

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "deployment"))
from aromin_stage import find_stage_change_column, header_key  # noqa: E402
from fix_stage_dates import name_hash  # noqa: E402


def canon_won(stage, fail):
    """همان canonFunnel: فقط «برنده بودن» لازم است"""
    if fail and str(fail).strip():
        return False
    s = re.sub(r"[\s‌]", "", str(stage or ""))
    if not s:
        return True
    if re.search("شکست|ازدست|باخت|لغو|ناموفق|مردود|ردشد", s):
        return False
    return bool(re.search("بستن|برنده|فروخته|فروش‌رفته", s))


def main(out, files):
    import openpyxl
    rows, meta = [], []
    for fi, path in enumerate(files):
        sha = hashlib.sha256(open(path, "rb").read()).hexdigest()
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        hdr = [str(h) if h is not None else "" for h in next(it, [])]
        # نام‌گذاریِ ستون‌های تکراری مثلِ SheetJS («شرکت»، «شرکت_1»)
        seen, keys = {}, []
        for h in hdr:
            keys.append(h if h not in seen else "%s_%d" % (h, seen[h]))
            seen[h] = seen.get(h, 0) + 1
        col, err = find_stage_change_column(keys)
        if err:
            meta.append(dict(name=os.path.basename(path), sha256=sha, sheet=ws.title, error=err))
            continue
        idx = {header_key(k): i for i, k in enumerate(keys)}
        ci, di = keys.index(col), idx.get(header_key("معامله"))
        si, fl = idx.get(header_key("مرحله")), idx.get(header_key("دلیل شکست"))
        names = [i for k, i in idx.items() if k in (header_key("شرکت"), header_key("مخاطب"))]
        n = 0
        for rn, r in enumerate(it, start=2):
            deal = r[di] if di is not None and di < len(r) else None
            if deal in (None, ""):
                continue
            n += 1
            v = r[ci] if ci < len(r) else None
            # فقط متن؛ سلولِ تاریخ/عددِ اکسل حدس زده نمی‌شود (raw_type ≠ str → نامعتبر، همان قاعدهٔ ایمپورت)
            rows.append(dict(deal=str(deal).strip(), raw=v if isinstance(v, str) else None, raw_type=type(v).__name__,
                             won=canon_won(r[si] if si is not None else "", r[fl] if fl is not None else ""),
                             name_h=[h for h in (name_hash(r[i]) for i in names if i < len(r)) if h],
                             file=fi, sheet=ws.title, row=rn))
        meta.append(dict(name=os.path.basename(path), sha256=sha, sheet=ws.title, column=col, rows=n))
    json.dump(dict(version=1, files=meta, rows=rows), open(out, "w", encoding="utf-8"), ensure_ascii=False)
    print(json.dumps(dict(files=meta, rows=len(rows)), ensure_ascii=False, indent=1))


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(2)
    main(sys.argv[1], sys.argv[2:])
